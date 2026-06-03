from datetime import datetime

from openpyxl import Workbook, load_workbook

from configuracion import COLUMNAS_BITACORA, COLUMNAS_ENTRADA


def crear_excel_entrada_si_no_existe(ruta_excel):
    ruta_excel.parent.mkdir(parents=True, exist_ok=True)

    if ruta_excel.exists():
        return

    libro = Workbook()
    hoja = libro.active
    hoja.title = "zips"
    hoja.append(COLUMNAS_ENTRADA)
    libro.save(ruta_excel)


def leer_fuentes_excel(ruta_excel):
    libro = load_workbook(ruta_excel)
    hoja = libro.active
    encabezados = [celda.value for celda in next(hoja.iter_rows(min_row=1, max_row=1))]
    posicion_nombre = encabezados.index("nombre") if "nombre" in encabezados else None
    posicion_url = encabezados.index("url")
    fuentes = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        url = fila[posicion_url]

        if not url:
            continue

        nombre = fila[posicion_nombre] if posicion_nombre is not None else ""

        fuentes.append({
            "nombre": nombre or "",
            "url": str(url).strip(),
        })

    return fuentes


def guardar_bitacora(ruta_excel, registros):
    ruta_excel.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "bitacora"
    hoja.append(COLUMNAS_BITACORA)

    for registro in registros:
        hoja.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            registro["nombre"],
            registro["origen"],
            registro["url_zip"],
            registro["archivo_zip"],
            registro["estado_descarga"],
            registro["estado_descompresion"],
            registro["carpeta_extraida"],
            registro["detalle"],
        ])

    libro.save(ruta_excel)
